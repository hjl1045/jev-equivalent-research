from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Subtitle',12),('Heading 1',19),('Heading 2',12)]:
 doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2']:doc.styles[name].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Auto claims classification  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Auto insurance document classification evaluation'
doc.core_properties.subject='Executive assessment of Jev Laya and GPT 5 6 Luna'
doc.core_properties.author=''
def p(text,bold=False):
 x=doc.add_paragraph();r=x.add_run(text);r.bold=bold;return x
def h(text):doc.add_heading(text,2)
def page(title):
 doc.add_page_break();doc.add_heading(title,1)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for cells,vals,idx in [(t.rows[0].cells,headers,0)]+[(t.add_row().cells,row,i+1) for i,row in enumerate(rows)]:
  for j,(c,value) in enumerate(zip(cells,vals)):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'17354A' if idx==0 else ('F1F5F7' if idx%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    m=OxmlElement('w:'+edge);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');margins.append(m)
   pr.append(margins)
   par=c.paragraphs[0];par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1);par.paragraph_format.line_spacing=1.03
   if j and widths[j]<1.1:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
   rr=par.add_run(str(value));rr.font.size=Pt(9.2);rr.bold=idx==0;rr.font.color.rgb=RGBColor.from_string('FFFFFF' if idx==0 else '000000')
  trpr=cells[0]._tc.getparent().get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t

doc.add_paragraph('Auto insurance document\nclassification evaluation','Title')
doc.add_paragraph('Executive assessment of Jev Laya and GPT 5 6 Luna','Subtitle')
p('22 September 2026  |  Feasibility pilot')
h('Evaluation summary')
p('The pilot compared Jev 1.13, default English Laya, and GPT-5.6 Luna on ten initial documents, eight additional boundary cases, and 12 length probes. Jev and Luna matched every author-assigned label. Laya made three errors on the initial set and four on the boundary set; chunked Laya classified all 12 long-input probes correctly. A later OpenJev demo run completed only two scored cases before quota exhaustion (page 16).')
table(['Evaluation','Measured result','Test coverage'],[
 ['Ten short documents','Jev 10/10\nLuna 10/10\nLaya 7/10','One example per category; 238–288 document tokens using the Laya tokenizer.'],
 ['Twelve length probes','Jev / Luna 12/12\nLaya chunks 12/12\nPrefix control 4/12','One police report in filler at about 1k, 4k, 5k, and 16k tokens; evidence at start, middle, or end.'],
 ['Eight boundary cases','Jev / Luna 8/8\nLaya 4/8','Related document purposes and one out-of-taxonomy collection letter; 92–120 document tokens.'],
 ['OpenRouter spend','$0.004184','30 Jev calls. Luna used Codex allowance; Laya ran locally. The spending cap was $0.05.']],[1.5,1.5,4])
h('Observed errors and input handling')
p('Laya confused neighboring purposes and assigned native confidence 0.9999 and 0.9943 to two incorrect boundary decisions. Its default configuration allowed 324 document tokens and silently clipped later text. It recognized only the four beginning-position length probes; chunking recovered all 12.')
h('Meaning of short and long in this study')
p('Short and long are test-set descriptions, not universal model thresholds. The short set contained 238–288 Laya tokens per document; length probes contained 1,027–16,388. The 5k probes contained 5,003–5,004 tokens: Jev and Luna each scored 3/3, default Laya 1/3, and chunked Laya 3/3. These lengths exceed default Laya’s document allowance.')
h('Evidence limits')
p('The synthetic cases use author-assigned labels and do not establish production accuracy or calibrated confidence. Long probes reuse one police report in neutral filler. The new boundary cases were frozen before inference and were designed after observing the initial errors; they are targeted diagnostics, not an independent random test set.')
p('Reading guide: original study on pages 2–10; native scores and boundary results on pages 11–12; complete boundary documents on pages 13–14; published capabilities and evidence on page 15.')

page('Study design and scope')
p('The evaluation addresses one operational question: given a standalone text document, which primary document type should receive it? The study compares a hosted decision model, an open-weight decision model, and a general-purpose LLM using the same category definitions.')
h('Evaluation population')
p('Ten fictional English documents represent ten categories: police report, demand letter, medical bill, medical record, repair estimate, rental invoice, witness statement, coverage letter, subrogation notice, and settlement release. An eleventh candidate label, other, provides a fallback, but the short-document set contains no positive example of that class.')
p('The short examples contain 238–288 Laya tokens and fit the available 324-token document budget. Several state their purpose explicitly and distinguish themselves from neighboring types. They are intentionally simple integration tests. All names, organizations, identifiers, and events are fictional.')
table(['Dimension','Implementation'],[
 ['Gold labels','Assigned by the document author from intended purpose; not independent expert adjudication.'],
 ['Jev','Pinned OpenRouter model typesafe/jev-1.13; returned version typesafe/jev-1.13-20260917. One document per request.'],
 ['Laya','convaiinnovations/laya English checkpoint; SDK 0.3.5; CPU with four threads; no fine-tuning.'],
 ['Luna','Requested gpt-5.6-luna via existing Codex ChatGPT login, reasoning none. One shuffled short batch; 12 independent long-input calls.'],
 ['Blinding','Gold labels and competing predictions excluded from inference inputs. Luna ran in an empty working directory and made no tool calls.'],
 ['Recorded evidence','Predictions, input hashes, raw responses, token usage, model/configuration metadata, timings, and Jev charges.']],[1.25,5.75])
h('Controlled long input probes')
p('The same police report was embedded in neutral administrative filler at the start, middle, or end of inputs near 1k, 4k, 5k, and 16k Laya tokens. The police report was selected because both Jev and Laya classified the short original correctly. Actual lengths ranged from 1,027 to 16,388 tokens. The 12 variants are correlated diagnostics, not independent cases.')
h('Fairness and exclusions')
p('Jev and Laya received identical document text and classification questions. Luna received the same taxonomy and documents in a different serving format. Its batched Codex run supports a correctness comparison, not a comparable per-document latency or dollar-cost result. Luna’s 12 long variants ran independently in empty working directories, without tools; CLI timings include Codex overhead. No OCR, images, production claims data, fine-tuning, or human workflow study was included.')

page('Classification results and operational implications')
table(['Document type','Jev','Luna','Laya prediction'],[
 ['Police report','Correct','Correct','Police report'],['Demand letter','Correct','Correct','Settlement release'],['Medical bill','Correct','Correct','Medical bill'],['Medical record','Correct','Correct','Medical record'],['Repair estimate','Correct','Correct','Repair estimate'],['Rental invoice','Correct','Correct','Rental invoice'],['Witness statement','Correct','Correct','Police report'],['Coverage letter','Correct','Correct','Coverage letter'],['Subrogation notice','Correct','Correct','Demand letter'],['Settlement release','Correct','Correct','Settlement release']],[2,1,1,3])
h('Three distinctions deserve focused testing')
p('Demand versus release. A demand proposes settlement and seeks a response; a release records agreement to give up specified claims. Laya chose release for the demand. In a deployed system, that error could place an unresolved negotiation into a completion workflow. This operational consequence is a risk scenario, not an observed loss.')
p('Witness account versus police report. Both describe a collision, but authorship and purpose differ. Laya treated the independent witness transcript as an officer report. Extraction and evidentiary review should preserve that distinction.')
p('Subrogation versus bodily injury demand. Both seek money. The recovery notice concerns reimbursement of an insurer payment already made; the attorney demand proposes resolution of an injury claim. Laya assigned the recovery request to the demand category.')
h('What can be concluded')
p('Jev and Luna passed all ten integration examples. Default Laya showed weaker separation of related types on this set. None of these findings demonstrates broad reliability: each positive category has one sample, the language is clear, and there is no unknown-document test. The study is too small to calibrate confidence or quantify category-specific error rates.')
h('Observed response times')
p('Jev’s median was 0.251 seconds per short document, including network round trip. Laya’s median was 0.389 seconds on a local CPU after loading and warmup. These observations compare deployed configurations, not intrinsic model speeds. No concurrency, throughput, or tail-latency benchmark was performed.')

page('Model mechanisms and sources of knowledge')
p('Jev and Laya are designed to return bounded decisions. The application supplies the document, a question, and candidate answers. A general LLM produces an answer through language generation; structured output can constrain that answer to a valid category. Both approaches still need to read the input. [1, 2, 3]')
table(['Aspect','Decision model','General LLM'],[
 ['Primary operation','Score supplied choices or evaluate a bounded condition.','Generate a label, structured object, extraction, or explanation.'],
 ['Output contract','Typed decision and probability distribution.','Constrained JSON or text, depending on configuration.'],
 ['Useful role','Focused routing, selection, and checks.','Ambiguity resolution, extraction, synthesis, and explanation.'],
 ['Important limitation','Valid structure and high confidence can accompany a wrong choice.','A valid JSON answer or self-reported confidence can also be wrong.']],[1.25,2.9,2.85])
h('Prior training supplies general language capability')
p('The models do not start with only the ten sample documents. Learned weights already encode language patterns and decision behavior. At inference time, category definitions tell the model how to apply that capability to this taxonomy. A model can recognize language about an invoice without knowing a particular insurer’s internal classification policy.')
p('TypeSafe describes Jev’s post-training as reinforcement learning for calibrated decisions, or RLCD, applied to pretrained language-model capabilities. Public documentation does not identify its complete training corpus or exact base architecture. It would be unsupported to claim that Jev was trained on insurance files or that it is a specific BERT model. Customer specialization is described through request context and criteria rather than per-account fine-tuning. [1, 2]')
p('Laya discloses a ModernBERT-large English encoder with a trained decision head. Its source scores candidate-option markers and normalizes the scores. The answer set can change at inference time. Its open weights permit local specialization, but capability on a specialized benchmark does not imply general zero-shot accuracy. [3, 4]')
h('Where proprietary knowledge belongs')
p('Supply document definitions, boundary examples, relevant claims procedures, and any case context needed for the decision. Use retrieval when the reference material is too large. The ten pilot documents were evaluation material; no training or prompt revision based on their results was performed.')
h('Confidence requires domain validation')
p('A class probability and a provider’s confidence field are not interchangeable. The Laya runtime warned about a shipped temperature for 11 or more options and clamped it. Neither model’s probabilities were calibrated against claims data here. Select review thresholds using a separate validation set and measure the errors that remain among accepted decisions.')

page('Long documents and evidence preservation')
table(['Model or configuration','Capacity','Practical interpretation'],[
 ['Jev through OpenRouter','32,000 context tokens','Reserve space for the question and candidate definitions.'],
 ['Jev direct API','64k total request; 32k state plus longest question','The larger request budget does not permit a 64k standalone document.'],
 ['Laya English default','512 total; 324 document tokens in this test','Installed SDK silently clips later document text.'],
 ['Laya alternate checkpoints','Published defaults of 1,024 total','Not tested here; do not assume the same behavior or quality.'],
 ['GPT-5.6 Luna','Published 1,050,000 context tokens','Tested through about 16k document tokens; published maximum not tested.']],[1.7,2,3.3])
p('Published limits are from provider documentation and endpoint metadata captured for this pilot. The Laya allowance was calculated with the installed sequence builder. Token counts are tokenizer-specific; pages and words are only approximations. [2, 3, 5, 6]')
table(['Length','Evidence positions','Jev','Luna','Laya prefix','Laya chunks'],[
 ['1k tokens','Start / middle / end','3/3','3/3','1/3','3/3'],['4k tokens','Start / middle / end','3/3','3/3','1/3','3/3'],['5k tokens','Start / middle / end','3/3','3/3','1/3','3/3'],['16k tokens','Start / middle / end','3/3','3/3','1/3','3/3']],[.85,1.65,.7,.9,1.45,1.45])
p('At 5,003–5,004 Laya tokens, default Laya retains about 6.5% of the input; chunk mode covers it in 18 windows. Jev and Luna received the full text. All three evidence positions were tested. This is a controlled retrieval diagnostic, not a representative sample of 5k-token claims documents.')
h('What the chunk experiment changed')
p('Laya’s prefix mode identified the report only when it appeared first. Chunk mode used 316-token windows with 32-token overlap, checked that re-encoded windows fit, and selected the strongest non-other winning chunk. It recovered all 12 variants. At 5k it took 7.73–7.74 seconds; Luna’s independent Codex runs took 3.22–3.46 seconds including CLI overhead, and Jev took 0.30–0.46 seconds. At 16k tokens it processed 58 chunks and took about 25.2–25.4 seconds per document; Jev processed those variants in about 0.35–0.44 seconds.')
p('This pooling rule is a diagnostic heuristic. It can favor one overconfident or irrelevant section, and its output is not a calibrated document probability. The repeated neutral filler was easier than a packet with real, conflicting exhibits.')
h('Production design implication')
p('Preserve document boundaries first. A demand packet may contain genuine bills, records, and police reports; its dominant page count may not reflect the covering letter’s purpose. Separate segment labels from packet labels, record the evidence location, and escalate conflicting sections. Reject silent truncation. Increasing Laya’s configured length requires validation of memory, supported positions, and quality rather than assuming the backbone limit guarantees reliable decisions. Jev also documents declining accuracy with irrelevant context. [7]')

page('Economics and operating model')
table(['Resource','Measured usage or cost','Interpretation'],[
 ['Jev short set','10 calls; $0.000342846','All paid calls used OpenRouter’s Decisions API.'],
 ['Jev long probes','12 calls; $0.003621240','Full-input requests at four length levels.'],
 ['Jev boundary cases','8 calls; $0.000219618','No retries; unchanged taxonomy and question.'],
 ['Total OpenRouter','$0.004183704 of $0.05 cap','About 0.418 cents across 30 calls.'],
 ['GPT-5.6 Luna','332,738 input; 471 output tokens','Codex allowance used; input includes Codex overhead. No dollar charge inferred.'],
 ['Local Laya','One English checkpoint; local CPU','No API fee, but compute, memory, setup, and maintenance have costs.']],[1.6,2.25,3.15])
h('Inference price is only one component')
p('Jev’s published input price was $0.042 per million tokens, with free output. At that rate, one million 1,000-token requests would cost about $42 in model input charges; one million 10,000-token requests would cost about $420. These are arithmetic illustrations assuming unchanged pricing and total billed input per request. They exclude OCR, storage, orchestration, retries, review labor, and integration. They are not a production forecast. [5]')
p('The pilot does not establish savings versus Luna: the Luna route used Codex allowance rather than a directly comparable API invoice. Laya is not costless merely because it is local. Measure cost per correctly routed document and reviewer time at the required accuracy, using the intended serving infrastructure.')
h('Recommended processing flow')
p('Receive text → identify document boundaries → check input quality and tokens → classify → apply validated review rules → route to the appropriate extraction workflow.',bold=True)
p('For a hosted pilot, evaluate Jev as the first classifier and Luna for ambiguous documents or extraction. If local deployment is a priority, test an adapted Laya configuration and explicit chunk aggregation against the same held-out set. A simple rules baseline should also be included for document types with dependable headings.')
h('Other candidate applications')
p('Focused decision models may fit message routing, intent detection, presence checks, selection from candidate extractions, relevance screening, and workflow branching. These uses were not benchmarked. General LLMs remain useful for open-ended writing and synthesis. Keep exact arithmetic and date comparisons in code, and distinguish a document’s mention of an attachment from proof that the attachment exists.')
h('Cost controls already implemented')
p('The Jev adapter checks pricing, restricts the provider, reserves budget before each request, records returned charges, and makes no automatic retries. Unknown billing retains its reservation. Four offline budget tests passed. This application guard does not change the account-wide spending limit.')

page('Implementation roadmap and decision gates')
p('The next phase should determine whether the classifier reduces intake work at an acceptable error rate. The following sequence and role assignments are proposals for planning, not committed staffing or delivery dates.')
table(['Stage','Work and proposed owner','Exit evidence'],[
 ['1  Define','Claims operations defines document unit, taxonomy, boundary examples, and review costs.','Approved category guide and adjudication process.'],
 ['2  Evaluate','Claims experts and evaluation lead label 20–50 cases per category plus unknown and mixed cases.','Frozen held-out set split by claim and template.'],
 ['3  Compare','Engineering runs Jev, Luna, Laya, and a rules baseline with recorded configurations.','Per-category errors, accepted-case accuracy, review rate, latency, and complete cost.'],
 ['4  Calibrate','Evaluation lead selects thresholds on separate validation data.','Documented tradeoff between error rate and manual review.'],
 ['5  Shadow','Operations and engineering compare suggestions with live handling before acting on them.','Stable quality across incoming sources and a tested fallback.'],
 ['6  Release','Business owner reviews evidence and limits initial automated scope.','Monitoring, rollback, version control, and ownership in place.']],[.9,3.35,2.75])
h('Acceptance criteria to agree before testing')
p('Set the permitted misrouting rate by category and business consequence, the required proportion of cases handled automatically, and the acceptable response time. A single aggregate accuracy target can hide mistakes in low-volume but important document types. Use confidence intervals and enough held-out observations to support the target; 20–50 examples per class is a starting point, not necessarily sufficient for rare-error claims.')
h('Required evaluation dimensions')
p('Report macro-F1 and per-category precision and recall, confusion patterns, unknown detection, malformed outputs, review rate, and correctness among automatically accepted cases. Measure full-pipeline latency under realistic concurrency, including retries and chunk aggregation. Preserve separate results for short documents, long documents, mixed packets, and damaged text.')
h('Controls that address the observed risks')
p('Block silent input loss; preserve page and segment boundaries; test misleading instructions embedded in documents; retain model and taxonomy versions; and capture evidence locations for review. Tune thresholds before release and revalidate them after a model change. Classification should route work; it should not silently become a coverage, liability, or payment decision.')

# Verbatim excerpts selected from the actual tested documents.
docs=[json.loads(s) for s in (ROOT/'data/documents.jsonl').read_text().splitlines()]
laya={r['id']:r['prediction'] for r in [json.loads(s) for s in (ROOT/'results/laya-smoke.jsonl').read_text().splitlines()]}
excerpts=[
'On April 18, 2026, at 16:42 I responded to the intersection of Finch Road and Harbor Avenue. The roadway was dry and visibility clear. Vehicle one, a blue sedan driven by Morgan Reed, showed front-end damage.',
'Considering the reported liability, treatment course, and disruption of daily activities, our client offers to resolve the bodily injury claim for $28,000 in exchange for a mutually acceptable release. Please provide your written response by September 24.',
'Total charges: $930.00. Payments posted: $100.00. Contractual adjustments: $80.00. Remaining balance: $750.00.\nPlease include the account identifier with payment. Questions about individual charges should be directed to the billing office.',
'Assessment: Improving cervical strain with residual mobility limitation. Progress is consistent with the treatment goals established at the initial visit. No new red-flag symptoms were reported during this encounter.',
'Replace rear bumper cover: part $420.00, body labor 1.8 hours at $75.00 = $135.00.\nReplace absorber: part $110.00, labor 0.5 hours at $75.00 = $37.50.\nRepair bracket alignment: labor 1.2 hours at $75.00 = $90.00.',
'Five rental days at $42.00 per day: $210.00. Local facility fee, five days at $3.00: $15.00. Tax: $18.00. Total: $243.00. Amount paid: $0.00. Balance due: $243.00.',
'Q: Please describe what you personally observed.\nA: The gray hatchback was waiting at the light. The blue sedan came up behind it and hit the rear. I heard braking just before the sound of the impact. I cannot estimate its speed.',
'We will investigate and handle the matter subject to a reservation of rights. Providing an adjuster, obtaining records, or discussing resolution does not waive any policy terms. No final coverage decision is made in this correspondence.',
'Bracken Insurance paid $3,860 under its insured\'s collision coverage for damage arising from the reported rear-end impact. We seek reimbursement of that payment and the insured\'s $500 deductible, for a total recovery request of $4,360.',
'In consideration of payment of $24,000, receipt and sufficiency of which are acknowledged subject to clearance of funds, Casey North agrees to release Morgan Reed and Alder Mutual from the bodily injury claims described in this agreement arising from the identified occurrence.'
]
for start in [0,5]:
 page('Sample document snapshots '+('one through five' if start==0 else 'six through ten'))
 p('Verbatim excerpts from the tested fictional documents. Expected labels reflect intended document purpose. Jev and Luna matched every expected label; Laya differences are shown explicitly.')
 for i in range(start,start+5):
  r=docs[i];assert excerpts[i] in r['text'],r['id']
  title=r['expected_label'].replace('_',' ').title()
  h(f"{i+1:02d}  {title}")
  x=p(excerpts[i]);x.paragraph_format.left_indent=Inches(.12);x.paragraph_format.space_after=Pt(5)
  p(f"Expected: {r['expected_label']}  |  ID: {r['id']}",True)
  predicted=laya[r['id']].replace('_',' ')
  x=p('Jev: correct  •  Luna: correct  •  Laya: '+('correct' if laya[r['id']]==r['expected_label'] else predicted+' — incorrect'))
  x.paragraph_format.space_after=Pt(11);x.runs[0].font.size=Pt(9)

page('Evidence boundaries and references')
h('Limits on interpretation')
p('The original study contains ten examples with author-assigned labels. Eight additional targeted boundary examples are reported separately on pages 11–14. The documents are clear, sometimes explicitly self-describing, and written in one style. The results do not estimate production accuracy, demonstrate statistical superiority, validate confidence calibration, or establish return on investment. Long probes reuse one document with repeated neutral filler. Chunk recovery on those probes does not validate classification of realistic mixed packets.')
p('Luna was tested in one initial batch, 12 independent long-input calls, and eight independent boundary calls through Codex, all with reasoning none and zero tool calls. The requested model id was recorded, but the captured stream does not expose a separately verified backend snapshot. Laya ran on CPU with SDK 0.3.5 and checkpoint revision 1c5edc17a7acd8701df6fc341c0d179f1c62c982. Jev returned version typesafe/jev-1.13-20260917. These differences constrain reproducibility and infrastructure comparisons.')
h('Evidence package')
p('The accompanying project retains the full examples in SYNTHETIC_DOCUMENTS.md and data/documents.jsonl; per-item results in results/comparison.md; length results in results/length-comparison.md; raw responses and the cost ledger in results/; and the benchmark and offline tests in scripts/. The snapshots on pages 8–9 are exact excerpts, not model-generated summaries.')
h('Provider and implementation references')
refs=[
('1','TypeSafe AI primer','Training objective and calibrated decisions','https://docs.typesafe.ai/introduction/machine-learning-primer'),
('2','TypeSafe models','Context limits and request-based customization','https://docs.typesafe.ai/models'),
('3','Laya model card','Checkpoint configurations and stated limitations','https://huggingface.co/convaiinnovations/laya'),
('4','Laya source repository','Encoder and decision-head implementation','https://github.com/NandhaKishorM/laya'),
('5','OpenRouter Jev 1.13','Model route and published pricing','https://openrouter.ai/typesafe/jev-1.13'),
('6','OpenAI GPT-5.6 Luna','Published context and structured-output capabilities','https://developers.openai.com/api/docs/models/gpt-5.6-luna'),
('7','TypeSafe Jev 1.13 limitations','Long-context distraction and other documented failure modes','https://docs.typesafe.ai/model-jaggedness/jev-1.13')]
for num,title,desc,url in refs:
 x=p(f'[{num}] {title}. {desc}.');x.paragraph_format.space_after=Pt(1);x.runs[0].font.size=Pt(9)
 x=p(url);x.paragraph_format.space_after=Pt(5);x.runs[0].font.size=Pt(8.5)
p('Documentation and endpoint metadata were checked on 22 September 2026. Published capacity and prices are not measurements of accuracy and can change.',False).runs[0].font.size=Pt(9)

# Native output appendix generated directly from retained responses.
score_rows=json.loads((ROOT/'results/confidence-details.json').read_text())
def score(source,id):return next(r for r in score_rows if r['source']==source and r['id']==id)
def compact(r):return f"{r['confidence']:.4f} / {r['selected_probability']:.4f}"
page('Native outputs and confidence scores')
p('Both Jev and Laya returned a selected choice, a probability for every candidate category, and a native confidence field. The tables show C / P: native confidence followed by the probability assigned to the selected label. Values are copied from retained responses; 1.0000 is a returned or rounded score, not proof of certainty.')
table(['Initial document','Jev C / P','Laya C / P','Laya outcome'],[[r['expected_label'].replace('_',' '),compact(score('jev-smoke.jsonl',r['id'])),compact(score('laya-smoke.jsonl',r['id'])),('Correct' if score('laya-smoke.jsonl',r['id'])['correct'] else score('laya-smoke.jsonl',r['id'])['prediction'].replace('_',' '))] for r in docs],[1.65,1.65,1.65,2.05])
h('What is distinctive in each output')
p('Jev also returned the dated model version, request identifier, provider, token usage, and billed cost. Laya also returned action.act_probability, a separate learned action-head output; it was 1.0000 for every initial and boundary example, including incorrect labels. It is not an independently measured probability that the category is correct.')
h('Confidence is different from selected class probability')
p('TypeSafe documents confidence as a statistic describing concentration of its probability distribution. Its public guide does not specify the exact formula. Laya’s installed source computes 1 minus normalized Shannon entropy: C = 1 − H(p) / ln(K), where K is the number of choices. These native confidence fields must not be assumed numerically interchangeable. Source: https://docs.typesafe.ai/confidence and installed laya/common.py, confidence_from_probs.')
p('The Laya SDK warned that the checkpoint’s temperature for 11 or more choices was outside its accepted range and clamped it. Its scores remain uncalibrated for this evaluation. Neither model was assessed on enough representative labeled claims to establish calibration. A native confidence of 0.99 is not an observed 99% success rate.')
h('Full distributions and long input scores')
p('results/confidence-details.json preserves all candidate probabilities, native confidence, top-two margin, action output, truncation status, and source file for every retained Jev and Laya result. The matching Markdown table is readable without code. For chunked Laya, the reported score belongs only to the selected chunk; no calibrated document-level confidence is produced. Luna’s label-only runs returned no comparable native confidence field.')

page('Boundary cases and observed confidence')
boundary=[json.loads(x) for x in (ROOT/'data/boundary_documents.jsonl').read_text().splitlines()]
p('Eight new examples deliberately distinguish related purposes under the unchanged 11-label taxonomy. They contain 92–120 Laya tokens, all below its 324-token allowance, so no truncation occurred. Gold labels were frozen before inference. Jev and Luna each scored 8/8; Laya scored 4/8. Luna ran each case independently through Codex without tools or access to the labels.')
table(['Case and expected','Jev choice and C / P','Laya choice and C / P'],[[r['id'].replace('boundary_','')+' '+r['expected_label'].replace('_',' '),score('jev-boundary.jsonl',r['id'])['prediction'].replace('_',' ')+'\n'+compact(score('jev-boundary.jsonl',r['id'])),score('laya-boundary.jsonl',r['id'])['prediction'].replace('_',' ')+'\n'+compact(score('laya-boundary.jsonl',r['id']))+('' if score('laya-boundary.jsonl',r['id'])['correct'] else ' — wrong')] for r in boundary],[2.1,2.4,2.5])
h('Scores on confusing inputs')
p('On the proposed release embedded in a demand, Laya selected settlement release with C = 0.9999 and P = 1.0000. On the officer-recorded witness interview, it selected police report with C = 0.9943 and P = 0.9985. Both labels were incorrect under the assigned primary-purpose labels. High native confidence did not eliminate these errors.')
p('Jev selected witness statement with C = 0.7800 and P = 0.8100; its alternative police report received 0.1900. On the repair-shop collection letter, Jev selected other with C = 0.4900 and P = 0.5400. This demonstrates a less concentrated answer on that example; it does not establish a calibrated review threshold.')
h('Interpretation boundaries')
p('The other label is correct for the collection letter under this specific taxonomy: demand_letter is restricted to attorney settlement requests for injury claims, and repair_estimate describes estimated repair costs. A broader business taxonomy could include a separate collection-notice category. Several examples explicitly clarify their purpose, making these controlled boundary checks easier than naturally ambiguous documents.')
p('These examples were constructed after the initial errors were observed. They test whether neighboring cues change the decision, not generalization to unseen claim populations. No prompt changes, fine-tuning, or repeated sampling were used. Complete text and expected labels follow; the original ten samples remain available in SYNTHETIC_DOCUMENTS.md.')

for start in [0,4]:
 page('Complete boundary documents '+('one through four' if start==0 else 'five through eight'))
 p('Complete model input text, not excerpts. Expected labels and neighboring categories below were held outside the inference input. All content is fictional.')
 for r in boundary[start:start+4]:
  h(r['id'].replace('boundary_','')+'  '+r['expected_label'].replace('_',' ').title())
  p(r['text'])
  x=p('Expected: '+r['expected_label']+'  |  Confusable with: '+r['confusable_label'],True);x.runs[0].font.size=Pt(9)


page('Published capabilities and supporting evidence')
p('Verified 22 September 2026. Published service limits, encoder capacity, and tested document budgets are distinct. No new inference was performed for this research update.')
table(['Model', 'Supported inputs', 'Published capacity', 'Tested scope'],[['Jev 1.13', 'Text: string, JSON object, or array of text values. No native image, audio or video. [S1]', 'Direct API: 64k total request; state + longest question limited to 32k. OpenRouter: 32,000 context. These are service limits, not a disclosed architectural maximum. [S1, S2]', 'Full text through about 16k document tokens. No image or binary ingestion tested.'], ['Laya English', 'Text and JSON serialized as text. Text encoder; no documented native image, audio or video path. [S3]', '512-token default sequence. ModernBERT backbone config lists 8,192 positions; that is not a validated Laya end-to-end limit. Multilingual default 1,024; its card states encoder capacity up to 8,192. [S3, S4]', '324 document tokens after this question. Longer documents used windows. No expanded-context checkpoint evaluation.'], ['GPT-5.6 Luna', 'Text and image input; text output. Audio and video unsupported in the model specification. [S5]', 'Published context: 1,050,000 tokens; maximum output: 128,000. Budget input, instructions and output within the context; Codex route availability/settings may differ. [S5]', 'Text only through about 16k document tokens. Images and published maximum context not tested.']],[.85,1.8,2.65,1.7])
p('PDF and DOCX are containers, not modalities. Jev and Laya require extracted text, including OCR for scans. Luna supports image input, but file upload and PDF conversion depend on the serving interface; native PDF ingestion was not verified here. Published context is not an accuracy guarantee, and a backbone limit does not validate the Laya decision head at that length.')
h('Primary source evidence')
p('S1  TypeSafe model limits and modalities',True).runs[0].font.size=Pt(9)
p('https://docs.typesafe.ai/models').runs[0].font.size=Pt(8)
p('S2  OpenRouter Jev route',True).runs[0].font.size=Pt(9)
p('https://openrouter.ai/typesafe/jev-1.13').runs[0].font.size=Pt(8)
p('S3  Laya model card and configurations',True).runs[0].font.size=Pt(9)
p('https://huggingface.co/convaiinnovations/laya').runs[0].font.size=Pt(8)
p('S4  ModernBERT encoder configuration',True).runs[0].font.size=Pt(9)
p('https://huggingface.co/answerdotai/ModernBERT-large/blob/main/config.json').runs[0].font.size=Pt(8)
p('S5  OpenAI Luna model specification',True).runs[0].font.size=Pt(9)
p('https://developers.openai.com/api/docs/models/gpt-5.6-luna').runs[0].font.size=Pt(8)

page('OpenJev additional evaluation and limitations')
p('Added 23 September 2026. The public community demo completed two scored initial documents correctly, then rejected the third request because the anonymous ZeroGPU quota was exhausted. OpenJev is partially tested; its results cannot be ranked against the completed three-model evaluation.')
table(['Test group','Completed','Observed outcome'],[['Initial ten documents','2 of 10','Police report and demand letter correct; remaining eight not run.'],['Eight boundary cases','0 of 8','Not run.'],['Twelve length probes','0 of 12','Not run.']],[2,1.3,3.7])
h('Actual route and score meaning')
p('The model-card-linked community Space serves openjev/openjev-FP8 dequantized to BF16 through Transformers on ZeroGPU. It returns probabilities from candidate-letter logits without the official helper choice temperature of 0.85. Its Gradio confidence fields are class probabilities, not a separately returned native confidence statistic. The full document text was preserved; shared instructions were prepended and label definitions appended to names, with commas replaced by semicolons for CSV transport.')
openjev_rows=[json.loads(x) for x in (ROOT/'results/openjev-demo.jsonl').read_text().splitlines()]
table(['Document','Prediction','Selected probability','Client seconds'],[[r['id'],r['prediction'],f"{r['selected_probability']:.4f}",f"{r['elapsed_s']:.2f}"] for r in openjev_rows],[1,2,2,2])
p('One preliminary connectivity pilot was excluded. Client timings include network, queue and GPU allocation; two observations do not establish typical speed. No paid inference calls were made. The demo source revision and method are recorded in results/openjev-demo.metadata.json; the quota rejection is preserved separately. No automatic retries were made.')
h('Published capabilities and deployment constraints')
p('The official model card describes text, JSON/DOM and screenshot input, one image per request, prompts up to 16,384 tokens, and 52 choices in one pass. Larger choice sets require several passes. The public demo exposes text and at most 26 labels. MLX builds are text-only. A 16k document plus instructions can exceed the official prompt budget; our longest probe has not been tested on OpenJev.')
p('The smallest published MLX build is about 15 GB; available local disk was 9.6 GiB. The weights are CC BY-NC 4.0, requiring attention to noncommercial licensing before any commercial deployment. OpenJev is independent of TypeSafe. An authenticated demo quota, suitable endpoint, or local checkpoint is required to finish this benchmark.')
h('Laya presentation correction')
p('Chunked Laya achieved 12/12 correct long-input classifications, including 3/3 at 5k tokens. The one-pager now leads with that result; the 4/12 prefix score is retained as a truncation control. Initial and boundary inputs fit one window, so their 7/10 and 4/8 results are separate from chunk recovery.')
p('Sources: https://huggingface.co/openjev/openjev and https://huggingface.co/spaces/chanoian/openjev-mlx-demo/blob/037a94430fdeda9956e97b8c0f30aea0652bced6/app.py').runs[0].font.size=Pt(8)
out=ROOT/'reports/2026-09-22-Auto-Claims-Classification-Executive-Report.docx';doc.save(out)
print(out)
